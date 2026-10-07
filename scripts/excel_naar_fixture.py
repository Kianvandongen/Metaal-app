"""Zet een calculatieblad (Excel, tabblad "Mat uittrekblad") om naar invoer voor de
calculatie-engine plus de verwachte uitkomsten uit Excel.

Gebruik: python3 -I scripts/excel_naar_fixture.py <blad.xlsx> <uit.json>

Bevat geen klantgegevens; de uitvoer wel, dus schrijf die buiten de repo.
"""
import json
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter as cl

SOORTEN = {
    "zagen": "zagen",
    "boren": "boren",
    "lassen": "lassen",
    "walsen": "walsen",
    "knippen": "knippen_buigen_snijden",
    "tekeningen": "tekeningen",
    "werkvoorbereiding": "werkvoorbereiding",
}
EXCEL = "overgenomen uit Excel (handmatig ingevuld)"


def is_formule(v):
    return isinstance(v, str) and v.startswith("=")


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def main(pad, uit):
    wf = openpyxl.load_workbook(pad, data_only=False)["Mat uittrekblad"]
    wv = openpyxl.load_workbook(pad, data_only=True)["Mat uittrekblad"]
    f = lambda c: wf[c].value
    v = lambda c: num(wv[c].value)

    # Bewerkingskolommen uit kopregel 3
    groepen = []
    for col in range(ci("J"), wf.max_column + 1):
        kop = wf.cell(3, col).value
        if not isinstance(kop, str):
            continue
        soort = next((s for k, s in SOORTEN.items() if k in kop.lower()), None)
        if not soort:
            continue
        per = col + 1 if cl(col) == "J" else col
        groepen.append((soort, cl(per), cl(per + 1)))

    tarieven = {s: 0.0 for s in SOORTEN.values()}
    for soort, _, tot in groepen:
        tarieven[soort] = v(f"{tot}26")

    posities = []
    for r in range(5, 25):
        aantal, lengte, kgm = v(f"C{r}"), v(f"D{r}"), v(f"G{r}")
        h = f(f"H{r}")
        p = {
            "nr": int(v(f"A{r}")) or r,
            "benaming": str(f(f"B{r}") or ""),
            "materiaal": str(f(f"E{r}") or ""),
            "aantal": aantal,
            "lengteMm": lengte,
            "kgPerM": kgm,
            "handelingen": v(f"J{r}"),
            "uurPerStuk": {},
            "urenHandmatig": {},
        }
        if h is not None and not (is_formule(h) and re.fullmatch(r"=G\d+\*\(F\d+/1000\)", h)):
            p["gewichtHandmatig"] = {"waarde": v(f"H{r}"), "reden": EXCEL + f": {h}"}
        for soort, per, tot in groepen:
            ps, tc = v(f"{per}{r}"), f(f"{tot}{r}")
            if r > 23:  # uren tellen in Excel t/m rij 23
                continue
            if ps:
                p["uurPerStuk"][soort] = ps
            if tc is not None and not is_formule(tc):
                p["urenHandmatig"][soort] = {"waarde": num(tc), "reden": EXCEL}
            elif tc is None and ps and p["handelingen"]:
                p["urenHandmatig"][soort] = {"waarde": 0, "reden": "leeg in Excel"}
        if r > 23:
            p["handelingen"] = 0
        posities.append(p)

    inkoop = []
    for r in range(33, 55):
        h, fcel, gcel = f(f"H{r}"), f(f"F{r}"), f(f"G{r}")
        oms = str(f(f"B{r}") or f(f"E{r}") or f"Inkoop rij {r}")
        if h is None:
            continue
        basis = "totaal_kg" if is_formule(fcel) and re.fullmatch(r"=(H25|F33)", fcel) else "stuk"
        regel = {"omschrijving": oms, "basis": basis, "aantal": v(f"F{r}"), "prijs": v(f"G{r}")}
        m = is_formule(gcel) and re.fullmatch(r"=([\d.]+)\*([\d.]+)", gcel)
        if m:  # bv. =220*1.1: inkoopprijs met opslag
            regel["prijs"] = float(m.group(1))
            regel["opslagPct"] = round((float(m.group(2)) - 1) * 100, 6)
        if not is_formule(h):
            regel["bedragHandmatig"] = {"waarde": v(f"H{r}"), "reden": EXCEL}
        inkoop.append(regel)

    kosten_col = next(c for c in ("Z", "W") if f(f"{c}32") == "Kosten")
    tot_col = cl(ci(kosten_col) + 1)
    m = re.fullmatch(rf"={tot_col}55\*([\d.]+)", str(f(f"{tot_col}56")))
    monteurs = float(m.group(1)) if m else 1
    montage = []
    for r in range(33, 55):
        if f(f"{kosten_col}{r}") is None or not f(f"L{r}"):
            continue
        montage.append({
            "omschrijving": str(f(f"L{r}")).strip(),
            "aantal": v(f"O{r}"),
            "tijd": v(f"P{r}"),
            "tarief": v(f"R{r}"),
            "perMonteur": True,  # Excel vermenigvuldigt alle montageregels met het aantal fte
        })

    marge = re.fullmatch(r"=H63\*([\d.]+)", str(f("H64")))
    bew_col = re.fullmatch(r"=([A-Z]+)28", str(f("H60"))).group(1)
    invoer = {
        "posities": posities,
        "kgPrijs": v("G26"),
        "zaagverliesPct": v("G27") * 100,
        "tarieven": tarieven,
        "inkoop": inkoop,
        "montage": montage,
        "aantalMonteurs": monteurs,
        "hulpmal": {"uren": v("F59"), "tarief": v("G59")},
        "margePct": float(marge.group(1)) * 100,
    }
    verwacht = {
        "totaalKg": v("H25"),
        "materiaal": v("H29"),
        "bewerkingen": v(f"{bew_col}28"),
        "inkoop": v("H55"),
        "montage": v("H61"),
        "hulpmal": v("H59"),
        "kostprijs": v("H63"),
        "marge": v("H64"),
        "verkoopprijs": v("H66"),
    }
    with open(uit, "w") as fh:
        json.dump({"bron": pad.split("/")[-1], "invoer": invoer, "verwacht": verwacht}, fh, indent=1)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
